from __future__ import annotations

import argparse
import os
import re
import warnings
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import psycopg2

from yearly_header_utils import HEADER_GROUPS, canonical_header, header_values, normalize_header


warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
    category=UserWarning,
)


def load_env_file(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue

        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]
        values[key] = value

    return values


def resolve_database_url() -> str | None:
    for key in ("DATABASE_URL", "SUPABASE_POOL_URL", "POSTGRES_URL"):
        value = os.environ.get(key)
        if value:
            return value

    repo_root = Path(__file__).resolve().parent
    for env_file in (repo_root / ".env.local", repo_root / ".env"):
        values = load_env_file(env_file)
        for key in ("DATABASE_URL", "SUPABASE_POOL_URL", "POSTGRES_URL"):
            value = values.get(key)
            if value:
                return value

    return None


def choose_data_sheet(workbook) -> str:
    best_sheet = workbook.sheetnames[0]
    best_score = -1
    for sheet_name in workbook.sheetnames:
        if sheet_name.strip().lower() in {"reference", "pivot"}:
            continue
        worksheet = workbook[sheet_name]
        first_row = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        score = sum(1 for value in first_row if value not in (None, ""))
        if normalize_header("Child Code") in {normalize_header(value) for value in first_row}:
            score += 100
        if score > best_score:
            best_score = score
            best_sheet = sheet_name
    return best_sheet


def infer_school_year_from_filename(filename: str) -> str:
    stem = Path(filename).stem
    match = re.search(r"(\d{4})[-_](\d{2,4})", stem)
    if not match:
        raise ValueError(f"Could not infer school year from filename: {filename}")
    start_year = int(match.group(1))
    end_part = match.group(2)
    if len(end_part) == 4:
        end_year = int(end_part)
    else:
        end_year = int(f"{start_year // 100}{end_part}")
        if end_year < start_year:
            end_year += 100
    return f"{start_year}-{end_year}"


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        text = " ".join(value.replace("\xa0", " ").split())
        return text or None
    text = str(value).strip()
    return text or None


def nonempty(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float, Decimal)):
        return value != 0
    text = clean_text(value)
    if text is None:
        return False
    normalized = text.strip().lower()
    return normalized in {"1", "y", "yes", "true", "t", "x", "checked", "present", "male", "female"}


def parse_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = clean_text(value)
    if text is None:
        return None
    text = text.replace(",", "")
    try:
        return int(float(text))
    except ValueError:
        return None


def parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return Decimal(int(value))
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = clean_text(value)
    if text is None:
        return None
    text = text.replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


_DATE_PATTERNS = [
    "%m/%d/%Y",
    "%m/%d/%y",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%b %d %Y",
    "%B %d %Y",
    "%d %b %Y",
    "%d %B %Y",
]


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if text is None:
        return None
    simplified = re.sub(r"[.,]", "", text.lower())
    simplified = re.sub(r"\s+", " ", simplified).strip()
    for pattern in _DATE_PATTERNS:
        try:
            return datetime.strptime(simplified, pattern).date()
        except ValueError:
            continue
    return None


def normalize_choice(value: Any) -> str | None:
    text = clean_text(value)
    return text


def normalize_lookup_label(value: Any) -> str | None:
    text = clean_text(value)
    return text


def is_row_empty(row: tuple[Any, ...]) -> bool:
    return not any(nonempty(value) for value in row)


def row_has_any(headers: list[object], row: tuple[Any, ...], *candidates: str) -> bool:
    for candidate in candidates:
        value = header_values(headers, row, candidate)
        if nonempty(value):
            return True
    return False


def value_for(headers: list[object], row: tuple[Any, ...], *candidates: str) -> Any:
    return header_values(headers, row, *candidates)


def build_insert_sql(table: str, data_columns: list[str], conflict_columns: list[str]) -> str:
    insert_columns_sql = ", ".join(data_columns)
    placeholders = ", ".join(["%s"] * len(data_columns))
    if conflict_columns:
        update_columns = [column for column in data_columns if column not in conflict_columns]
        if update_columns:
            update_sql = ", ".join(
                f"{column} = coalesce(excluded.{column}, {table}.{column})" for column in update_columns
            )
            return (
                f"insert into {table} ({insert_columns_sql}) values ({placeholders}) "
                f"on conflict ({', '.join(conflict_columns)}) do update set {update_sql} returning id"
            )
        return (
            f"insert into {table} ({insert_columns_sql}) values ({placeholders}) "
            f"on conflict ({', '.join(conflict_columns)}) do nothing returning id"
        )
    return f"insert into {table} ({insert_columns_sql}) values ({placeholders}) returning id"


@dataclass
class ImportStats:
    rows_seen: int = 0
    rows_skipped: int = 0
    students: int = 0
    enrollments: int = 0
    household_rows: int = 0
    housing_rows: int = 0
    welfare_rows: int = 0
    academic_rows: int = 0
    assistance_rows: int = 0
    warnings: int = 0


class YearlyImporter:
    def __init__(self, connection, dry_run: bool = False, verbose: bool = False):
        self.connection = connection
        self.cursor = connection.cursor() if connection is not None else None
        self.dry_run = dry_run
        self.verbose = verbose
        self.stats = ImportStats()
        self.lookup_cache: dict[tuple[Any, ...], Any] = {}

    def close(self) -> None:
        if self.cursor is not None:
            self.cursor.close()
            self.cursor = None

    def maybe_commit(self) -> None:
        if self.dry_run or self.connection is None:
            return
        if self.stats.rows_seen % 250 == 0:
            self.connection.commit()
            if self.verbose and self.stats.rows_seen:
                print(f"  committed batch at {self.stats.rows_seen} rows")

    def execute_with_retry(self, sql: str, params: tuple[Any, ...] | list[Any] | None = None) -> Any:
        if self.dry_run or self.cursor is None:
            return None
        last_error: Exception | None = None
        for attempt in range(3):
            try:
                self.cursor.execute(sql, params or ())
                return self.cursor.fetchone()
            except psycopg2.OperationalError as exc:
                last_error = exc
                if attempt == 2:
                    break
                try:
                    self.connection.rollback()
                except Exception:
                    pass
                self.connection.reset()
        raise last_error

    def warn(self, message: str) -> None:
        self.stats.warnings += 1
        print(f"  WARN: {message}")

    def fetch_lookup(self, key: tuple[Any, ...], table: str, insert_columns: list[str], conflict_columns: list[str], values: list[Any]) -> Any:
        if key in self.lookup_cache:
            return self.lookup_cache[key]
        if self.dry_run:
            synthetic = f"dry-run:{table}:{len(self.lookup_cache) + 1}"
            self.lookup_cache[key] = synthetic
            return synthetic
        sql = build_insert_sql(table, insert_columns, conflict_columns)
        self.cursor.execute(sql, values)
        row = self.cursor.fetchone()
        if row is None:
            self.cursor.execute(
                f"select id from {table} where {', '.join(conflict_columns)} = %s",
                (values[conflict_columns.index(conflict_columns[0])],),
            )
            row = self.cursor.fetchone()
        lookup_id = row[0]
        self.lookup_cache[key] = lookup_id
        return lookup_id

    def upsert_lookup_single(self, table: str, column: str, value: Any) -> Any:
        label = normalize_lookup_label(value)
        if label is None:
            return None
        key = (table, column, normalize_header(label))
        if key in self.lookup_cache:
            return self.lookup_cache[key]
        if self.dry_run:
            synthetic = f"dry-run:{table}:{len(self.lookup_cache) + 1}"
            self.lookup_cache[key] = synthetic
            return synthetic

        self.cursor.execute(f"select id from {table} where {column} = %s", (label,))
        row = self.cursor.fetchone()
        if row is not None:
            lookup_id = row[0]
            self.lookup_cache[key] = lookup_id
            return lookup_id

        try:
            self.cursor.execute(f"insert into {table} ({column}) values (%s) returning id", (label,))
            row = self.cursor.fetchone()
            lookup_id = row[0] if row else None
        except psycopg2.IntegrityError:
            self.connection.rollback()
            self.cursor.execute(f"select id from {table} where {column} = %s", (label,))
            row = self.cursor.fetchone()
            lookup_id = row[0] if row else None

        self.lookup_cache[key] = lookup_id
        return lookup_id

    def upsert_lookup_city(self, province_id: Any, value: Any) -> Any:
        city = normalize_lookup_label(value)
        if city is None:
            return None
        key = ("lookup_cities", province_id, normalize_header(city))
        if key in self.lookup_cache:
            return self.lookup_cache[key]
        if self.dry_run:
            synthetic = f"dry-run:lookup_cities:{len(self.lookup_cache) + 1}"
            self.lookup_cache[key] = synthetic
            return synthetic
        self.cursor.execute(
            "select id from lookup_cities where province_id is not distinct from %s and name = %s",
            (province_id, city),
        )
        row = self.cursor.fetchone()
        if row is not None:
            city_id = row[0]
            self.lookup_cache[key] = city_id
            return city_id

        try:
            self.cursor.execute(
                "insert into lookup_cities (province_id, name) values (%s, %s) returning id",
                (province_id, city),
            )
            row = self.cursor.fetchone()
            city_id = row[0] if row else None
        except psycopg2.IntegrityError:
            self.connection.rollback()
            self.cursor.execute(
                "select id from lookup_cities where province_id is not distinct from %s and name = %s",
                (province_id, city),
            )
            row = self.cursor.fetchone()
            city_id = row[0] if row else None

        self.lookup_cache[key] = city_id
        return city_id

    def upsert_student(self, row: tuple[Any, ...], headers: list[object]) -> Any:
        child_code = clean_text(value_for(headers, row, "Child Code"))
        if not child_code:
            return None

        province_id = self.upsert_lookup_single("lookup_provinces", "name", value_for(headers, row, "Province"))
        city_id = self.upsert_lookup_city(province_id, value_for(headers, row, "City/ Municipality", "edited_City/ Municipality", "Municipality/ City"))

        data = {
            "child_code": child_code,
            "surname": clean_text(value_for(headers, row, "*Child's Surname", "Child's Surname")),
            "first_name": clean_text(value_for(headers, row, "*Child's First Name", "Child's First Name")),
            "middle_initial": clean_text(value_for(headers, row, "*Middle Initial", "Middle Initial")),
            "sex": clean_text(value_for(headers, row, "* Sex/ Gender", "Sex/ Gender", "Sex")),
            "date_of_birth": parse_date(value_for(headers, row, "*Date of Birth", "Date of Birth")),
            "province_id": province_id,
            "city_id": city_id,
            "barangay": clean_text(value_for(headers, row, "*Barangay", "Barangay")),
            "sitio_phase": clean_text(value_for(headers, row, "*Sitio/ Phase", "*Sitio Phase", "Sitio/ Phase", "Sitio Phase")),
            "complete_address": clean_text(value_for(headers, row, "*Complete Address", "Complete Address")),
        }

        sex = data["sex"]
        if sex:
            normalized_sex = sex.strip().lower()
            if normalized_sex in {"m", "male"}:
                data["sex"] = "Male"
            elif normalized_sex in {"f", "female"}:
                data["sex"] = "Female"

        if self.dry_run:
            self.stats.students += 1
            return f"dry-run:student:{child_code}"

        sql = build_insert_sql(
            "students",
            [
                "child_code",
                "surname",
                "first_name",
                "middle_initial",
                "sex",
                "date_of_birth",
                "province_id",
                "city_id",
                "barangay",
                "sitio_phase",
                "complete_address",
            ],
            ["child_code"],
        )
        student_id = self.execute_with_retry(
            sql,
            (
                data["child_code"],
                data["surname"],
                data["first_name"],
                data["middle_initial"],
                data["sex"],
                data["date_of_birth"],
                data["province_id"],
                data["city_id"],
                data["barangay"],
                data["sitio_phase"],
                data["complete_address"],
            ),
        )
        if student_id is None:
            raise RuntimeError("Failed to create student record")
        student_id = student_id[0]
        self.stats.students += 1
        return student_id

    def upsert_enrollment(self, student_id: Any, school_year: str, row: tuple[Any, ...], headers: list[object]) -> Any:
        grade_level = clean_text(value_for(headers, row, "*Grade/ Year Level", "*Grade/ Year Level SY2020-21", "Grade/ Year Level", "Grade/ Year Level SY2020-21"))
        shs_track = clean_text(value_for(headers, row, "Track for Senior High", "Track/Strand (for Senior High Students)"))
        course_in_college = clean_text(value_for(headers, row, "Course in College", "Course (for College Students)"))
        school_name = clean_text(value_for(headers, row, "*Name of School", "Name of School"))
        classification = clean_text(value_for(headers, row, "*Child's Classification upon Admission", "Child's Classification upon Admission"))
        cnsp = clean_text(value_for(headers, row, "Children in Need of Special Protection ( CNSP Cluster)", "Children in Need of Special Protection (CNSP Cluster)"))
        dropout_reason = clean_text(value_for(headers, row, "Reason for dropping out", "Reason for dropping out: Others", "Reason/s for Dropping out ( Remarks from Teachers of POs/ SDWs"))
        education_status = clean_text(value_for(headers, row, "Child's Education  Status  (as ERDA Beneficiary)", "Education Status"))
        location_for_reporting = clean_text(value_for(headers, row, "Location for Reporting"))
        remarks = clean_text(value_for(headers, row, "*REMARKS", "REMARKS", "Remarks"))
        batch_number = clean_text(value_for(headers, row, "Batch Number"))
        intake_no = parse_int(value_for(headers, row, "No.", "No. "))
        intake_date = parse_date(value_for(headers, row, "Date of Intake Interview"))

        grade_level_id = self.upsert_lookup_single("lookup_grade_levels", "name", grade_level)
        shs_track_id = self.upsert_lookup_single("lookup_shs_tracks", "name", shs_track)
        classification_id = self.upsert_lookup_single("lookup_classifications", "name", classification)
        cnsp_cluster_id = self.upsert_lookup_single("lookup_cnsp_clusters", "name", cnsp)

        if self.dry_run:
            self.stats.enrollments += 1
            return f"dry-run:enrollment:{student_id}:{school_year}"

        sql = build_insert_sql(
            "enrollment_records",
            [
                "student_id",
                "school_year",
                "batch_number",
                "intake_no",
                "date_of_intake_interview",
                "grade_level_id",
                "shs_track_id",
                "course_in_college",
                "school_name",
                "classification_id",
                "dropout_reason",
                "cnsp_cluster_id",
                "education_status",
                "location_for_reporting",
                "remarks",
            ],
            ["student_id", "school_year"],
        )
        enrollment_id = self.execute_with_retry(
            sql,
            (
                student_id,
                school_year,
                batch_number,
                intake_no,
                intake_date,
                grade_level_id,
                shs_track_id,
                course_in_college,
                school_name,
                classification_id,
                dropout_reason,
                cnsp_cluster_id,
                education_status,
                location_for_reporting,
                remarks,
            ),
        )
        if enrollment_id is None:
            raise RuntimeError("Failed to create enrollment record")
        enrollment_id = enrollment_id[0]
        self.stats.enrollments += 1
        return enrollment_id

    def upsert_household(self, enrollment_id: Any, row: tuple[Any, ...], headers: list[object]) -> None:
        hh_head_surname = clean_text(value_for(headers, row, "HH Head: Surname", "HH Head: Surname ( Mother)", "HH Head: Surname (Mother)", "HH Head: Surname(father)", "HH Head: Surname(mother)"))
        hh_head_first_name = clean_text(value_for(headers, row, "HH Head: First Name", "HH Head: First Name ( Mother)", "HH Head: First Name (Mother)", "HH Head: First Name(father)", "HH Head: First Name(mother)"))
        hh_head_mi = clean_text(value_for(headers, row, "HH Head: MI"))
        occupation = clean_text(value_for(headers, row, "Occupation of HH Head"))
        pci_value = value_for(headers, row, "Family Per Capita Income", "pci below 800")
        pci_below_800 = parse_bool(value_for(headers, row, "pci below 800"))
        income_label = normalize_lookup_label(pci_value)
        income_bracket_id = self.upsert_lookup_single("lookup_income_brackets", "label", income_label) if income_label else None

        if not any(nonempty(value) for value in [hh_head_surname, hh_head_first_name, hh_head_mi, occupation, income_label, pci_below_800]):
            return

        if self.dry_run:
            self.stats.household_rows += 1
            return

        sql = build_insert_sql(
            "household_snapshots",
            ["enrollment_record_id", "hh_head_surname", "hh_head_first_name", "hh_head_mi", "occupation_of_hh_head", "income_bracket_id", "pci_below_800"],
            ["enrollment_record_id"],
        )
        self.cursor.execute(
            sql,
            (enrollment_id, hh_head_surname, hh_head_first_name, hh_head_mi, occupation, income_bracket_id, pci_below_800),
        )
        self.stats.household_rows += 1

    def upsert_housing(self, enrollment_id: Any, row: tuple[Any, ...], headers: list[object]) -> None:
        house_ownership = clean_text(value_for(headers, row, "HC: House Ownership", "HC: House Ownership: If others, specify", "HC: House Ownership: Others"))
        house_ownership_other = clean_text(value_for(headers, row, "HC: House Ownership: Others", "HC: House Ownership: If others, specify"))
        size_of_dwelling = clean_text(value_for(headers, row, "Size of Dwelling"))
        materials = clean_text(value_for(headers, row, "Materials", "Materials: Others", "Materials:Others", "Materials: If others, specify"))
        materials_other = clean_text(value_for(headers, row, "Materials: Others", "Materials:Others", "Materials: If others, specify"))
        water_supply = clean_text(value_for(headers, row, "Water Supply", "Water Supply: Others", "Water Supply: If others, specify"))
        water_supply_other = clean_text(value_for(headers, row, "Water Supply: Others", "Water Supply: If others, specify"))
        lighting = clean_text(value_for(headers, row, "Lighting Facilities", "Lighting Facilities: Others", "Lighting Facilities: If others, specify"))
        lighting_other = clean_text(value_for(headers, row, "Lighting Facilities: Others", "Lighting Facilities: If others, specify"))
        toilet = clean_text(value_for(headers, row, "Toilet Facility", "Toilet Facility: Others", "Toilet Facility: If others, specify"))
        toilet_other = clean_text(value_for(headers, row, "Toilet Facility: Others", "Toilet Facility: If others, specify"))
        furnitures = clean_text(value_for(headers, row, "Furniture/s"))

        if not any(nonempty(value) for value in [house_ownership, house_ownership_other, size_of_dwelling, materials, materials_other, water_supply, water_supply_other, lighting, lighting_other, toilet, toilet_other, furnitures]):
            return

        house_ownership_id = self.upsert_lookup_single("lookup_house_ownership_types", "name", house_ownership)
        materials_id = self.upsert_lookup_single("lookup_dwelling_materials", "name", materials)
        water_supply_id = self.upsert_lookup_single("lookup_water_supply_types", "name", water_supply)
        lighting_id = self.upsert_lookup_single("lookup_lighting_facility_types", "name", lighting)
        toilet_id = self.upsert_lookup_single("lookup_toilet_facility_types", "name", toilet)

        if self.dry_run:
            self.stats.housing_rows += 1
            return

        sql = build_insert_sql(
            "housing_conditions",
            [
                "enrollment_record_id",
                "house_ownership_id",
                "house_ownership_other",
                "size_of_dwelling",
                "materials_id",
                "materials_other",
                "water_supply_id",
                "water_supply_other",
                "lighting_facility_id",
                "lighting_facility_other",
                "toilet_facility_id",
                "toilet_facility_other",
                "furnitures",
            ],
            ["enrollment_record_id"],
        )
        self.cursor.execute(
            sql,
            (
                enrollment_id,
                house_ownership_id,
                house_ownership_other,
                size_of_dwelling,
                materials_id,
                materials_other,
                water_supply_id,
                water_supply_other,
                lighting_id,
                lighting_other,
                toilet_id,
                toilet_other,
                furnitures,
            ),
        )
        self.stats.housing_rows += 1

    def upsert_welfare(self, enrollment_id: Any, row: tuple[Any, ...], headers: list[object]) -> None:
        presently_served = parse_bool(value_for(headers, row, "Presently served by welfare agency"))
        welfare_agency_name = clean_text(value_for(headers, row, "Name of  Welfare Agency", "Name of Welfare Agency"))
        type_of_service = clean_text(value_for(headers, row, "Type of service Availed", "Type of service availed", "Type of Service Availed"))
        is_org_member = parse_bool(value_for(headers, row, "Membership in any organization", "Membership to school organization/s"))
        organization_type = clean_text(value_for(headers, row, "Type of Organization"))
        education_scheme = clean_text(value_for(headers, row, "Type of Scheme/ Implementation ", "Type of Scheme/ Implementation"))
        social_protection_program = clean_text(value_for(headers, row, "Social Protection  Program Affiliation ", "Social Protection Program Affiliation", "Social Protection: Others"))
        social_protection_other = clean_text(value_for(headers, row, "Social Protection: Others"))
        tie_up_partner = clean_text(value_for(headers, row, "Name of Tie-up Partner"))
        staff_in_charge = clean_text(value_for(headers, row, "Staff In- Charge", "Staff In-Charge"))
        funder_name = clean_text(value_for(headers, row, "Name of Funder     [for Finance Section]", "Name of Funder"))

        if not any(nonempty(value) for value in [presently_served, welfare_agency_name, type_of_service, is_org_member, organization_type, education_scheme, social_protection_program, social_protection_other, tie_up_partner, staff_in_charge, funder_name]):
            return

        education_scheme_id = self.upsert_lookup_single("lookup_education_schemes", "name", education_scheme)

        if self.dry_run:
            self.stats.welfare_rows += 1
            return

        sql = build_insert_sql(
            "welfare_services",
            [
                "enrollment_record_id",
                "presently_served_by_welfare_agency",
                "welfare_agency_name",
                "type_of_service_availed",
                "is_organization_member",
                "organization_type",
                "education_scheme_id",
                "social_protection_program",
                "social_protection_other",
                "tie_up_partner_name",
                "staff_in_charge",
                "funder_name",
            ],
            ["enrollment_record_id"],
        )
        self.cursor.execute(
            sql,
            (
                enrollment_id,
                presently_served,
                welfare_agency_name,
                type_of_service,
                is_org_member,
                organization_type,
                education_scheme_id,
                social_protection_program,
                social_protection_other,
                tie_up_partner,
                staff_in_charge,
                funder_name,
            ),
        )
        self.stats.welfare_rows += 1

    def upsert_academic(self, enrollment_id: Any, row: tuple[Any, ...], headers: list[object]) -> None:
        academic = {
            "nc_coc": clean_text(value_for(headers, row, "NC/COC")),
            "enrolled_and_assisted_current_sy": parse_bool(value_for(headers, row, "Enrolled and Assisted in the Current SY")),
            "education_status_as_beneficiary": clean_text(value_for(headers, row, "Child's Education  Status  (as ERDA Beneficiary)", "Education Status")),
            "reason_for_being_inactive": clean_text(value_for(headers, row, "Reason for being Inactive", "Others Please Specify/Reason/s why child lost interest")),
            "graduate_remarks": clean_text(value_for(headers, row, "Graduate_Remarks")),
            "dropped_out_past_4_months": parse_bool(value_for(headers, row, "Did the child drop out from school for the past four months")),
            "dropout_reason_teacher_remarks": clean_text(value_for(headers, row, "Reason/s for Dropping out ( Remarks from Teachers of POs/ SDWs")),
            "dropout_reason_other": clean_text(value_for(headers, row, "Reason for dropping out: Others")),
            "dropout_month": clean_text(value_for(headers, row, "Specific Month of dropping out from School ( Please indicate Month and Year)")),
            "school_awards": clean_text(value_for(headers, row, "School Awards received by the child        non-academic/ academic-end of SY")),
            "school_activities_attended": clean_text(value_for(headers, row, "School Activities Attended  by the Child")),
            "school_org_membership": clean_text(value_for(headers, row, "Membership to school organization/s", "Membership in any organization")),
            "capacity_building_activities": clean_text(value_for(headers, row, "Capacity Building  Activities Attended by the child within the community")),
            "math_grade_1st_period": parse_decimal(value_for(headers, row, "Math Grade    ( 1st Grading Period)       For Children attending dear and mathemagica")),
            "general_average": parse_decimal(value_for(headers, row, "General Average       For Children attending dear and mathemagica")),
            "reading_skills_pre_assessment": clean_text(value_for(headers, row, "Level of Reading Skills (Pre- assessment) for Children Attending dear and mathemagica)")),
            "reading_skills_post_assessment": clean_text(value_for(headers, row, "Level of Reading Skills (Post -assessment)")),
        }

        if not any(nonempty(value) for value in academic.values()):
            return

        if self.dry_run:
            self.stats.academic_rows += 1
            return

        sql = build_insert_sql(
            "academic_performance",
            ["enrollment_record_id", *academic.keys()],
            ["enrollment_record_id"],
        )
        self.cursor.execute(sql, (enrollment_id, *academic.values()))
        self.stats.academic_rows += 1

    def insert_assistance(self, enrollment_id: Any, row: tuple[Any, ...], headers: list[object]) -> None:
        records: list[dict[str, Any]] = []

        social_type = clean_text(value_for(headers, row, "Assistance  availed or received: (Social Protection)"))
        social_date = parse_date(value_for(headers, row, "When did the child  received the Assistance   ( Specify Month)"))
        visit_date = parse_date(value_for(headers, row, " Date of Visit (format: month, day, year)"))
        remarks = clean_text(value_for(headers, row, "Remarks", "*REMARKS"))
        if any(nonempty(value) for value in [social_type, social_date, visit_date, remarks]):
            records.append(
                {
                    "assistance_category": "Social Protection",
                    "assistance_type": social_type,
                    "date_received": social_date,
                    "visit_date": visit_date,
                    "remarks": remarks,
                }
            )

        edu_type = clean_text(value_for(headers, row, "Type of assistance received by the child from ERDA ( Educational Assistance)"))
        edu_date = parse_date(value_for(headers, row, "When did the child  availed or received the Assistance   ( Specify Month & year)"))
        if any(nonempty(value) for value in [edu_type, edu_date, remarks]):
            records.append(
                {
                    "assistance_category": "Educational Assistance",
                    "assistance_type": edu_type,
                    "date_received": edu_date,
                    "visit_date": visit_date,
                    "remarks": remarks,
                }
            )

        if not records:
            return

        if self.dry_run:
            self.stats.assistance_rows += len(records)
            return

        self.cursor.execute("delete from assistance_records where enrollment_record_id = %s", (enrollment_id,))
        sql = "insert into assistance_records (enrollment_record_id, assistance_category, assistance_type, date_received, visit_date, remarks) values (%s, %s, %s, %s, %s, %s)"
        for record in records:
            self.cursor.execute(
                sql,
                (
                    enrollment_id,
                    record["assistance_category"],
                    record["assistance_type"],
                    record["date_received"],
                    record["visit_date"],
                    record["remarks"],
                ),
            )
        self.stats.assistance_rows += len(records)

    def import_workbook(self, workbook_path: Path, school_year: str) -> ImportStats:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_name = choose_data_sheet(workbook)
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]

        if self.verbose:
            print(f"Using sheet {sheet_name!r} for {workbook_path.name}")

        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if row is None or is_row_empty(row):
                self.stats.rows_skipped += 1
                continue
            self.stats.rows_seen += 1
            self.maybe_commit()

            child_code = clean_text(value_for(headers, row, "Child Code"))
            if not child_code:
                self.warn(f"{workbook_path.name}:{sheet_name}:{row_number} missing Child Code; row skipped")
                self.stats.rows_skipped += 1
                continue

            student_id = self.upsert_student(row, headers)
            if student_id is None:
                self.warn(f"{workbook_path.name}:{sheet_name}:{row_number} could not create student for child_code={child_code}")
                self.stats.rows_skipped += 1
                continue

            enrollment_id = self.upsert_enrollment(student_id, school_year, row, headers)
            self.upsert_household(enrollment_id, row, headers)
            self.upsert_housing(enrollment_id, row, headers)
            self.upsert_welfare(enrollment_id, row, headers)
            self.upsert_academic(enrollment_id, row, headers)
            self.insert_assistance(enrollment_id, row, headers)

        return self.stats


def run_import(workbook_path: Path, school_year: str | None = None, dry_run: bool = False, verbose: bool = False) -> ImportStats:
    workbook_path = workbook_path.resolve()
    if school_year is None:
        school_year = infer_school_year_from_filename(workbook_path.name)

    database_url = resolve_database_url()
    if not dry_run and not database_url:
        raise RuntimeError("A database connection string is required for a real migration run (DATABASE_URL or SUPABASE_POOL_URL)")

    if dry_run:
        connection = None
        importer = YearlyImporter(connection, dry_run=True, verbose=verbose)
        try:
            return importer.import_workbook(workbook_path, school_year)
        finally:
            importer.close()

    connection = psycopg2.connect(database_url)
    try:
        connection.autocommit = False
        with connection.cursor() as cursor:
            cursor.execute("set statement_timeout = '10min'")
            cursor.execute("set session characteristics as transaction read write")
        importer = YearlyImporter(connection, dry_run=False, verbose=verbose)
        try:
            stats = importer.import_workbook(workbook_path, school_year)
            connection.commit()
            return stats
        except Exception:
            try:
                connection.rollback()
            except Exception:
                pass
            raise
        finally:
            importer.close()
    finally:
        try:
            connection.close()
        except Exception:
            pass


def format_stats(path: Path, school_year: str, stats: ImportStats) -> str:
    return (
        f"{path.name} [{school_year}] rows_seen={stats.rows_seen} rows_skipped={stats.rows_skipped} "
        f"students={stats.students} enrollments={stats.enrollments} household={stats.household_rows} "
        f"housing={stats.housing_rows} welfare={stats.welfare_rows} academic={stats.academic_rows} "
        f"assistance={stats.assistance_rows} warnings={stats.warnings}"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Import a single ERDA yearly intake workbook into Supabase/Postgres.")
    parser.add_argument("workbook", help="Path to the .xlsx file")
    parser.add_argument("school_year", nargs="?", help="School year label to store, e.g. 2017-2018")
    parser.add_argument("--dry-run", action="store_true", help="Parse the workbook and print stats without writing to the database")
    parser.add_argument("--verbose", action="store_true", help="Print extra progress information")
    args = parser.parse_args()

    stats = run_import(Path(args.workbook), args.school_year, dry_run=args.dry_run, verbose=args.verbose)
    school_year = args.school_year or infer_school_year_from_filename(args.workbook)
    print(format_stats(Path(args.workbook), school_year, stats))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
