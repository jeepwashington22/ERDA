from openpyxl import load_workbook
from pathlib import Path
path = Path('YearlyRecords/SY 2017-18.xlsx')
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
print('sheet', wb.sheetnames[0])
print('rows', ws.max_row)
