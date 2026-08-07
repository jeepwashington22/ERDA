from __future__ import annotations

import argparse
from difflib import get_close_matches
from pathlib import Path
import warnings

from openpyxl import load_workbook

from yearly_header_utils import HEADER_LOOKUP, canonical_header, normalize_header


warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
    category=UserWarning,
)


def choose_data_sheet(workbook) -> str:
    best_sheet = workbook.sheetnames[0]
    best_score = -1
    for sheet_name in workbook.sheetnames:
        lower_name = sheet_name.strip().lower()
        if lower_name in {"reference", "pivot"}:
            continue
        worksheet = workbook[sheet_name]
        first_row = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        score = sum(1 for value in first_row if value not in (None, ""))
        if "child code" in {normalize_header(value) for value in first_row}:
            score += 100
        if score > best_score:
            best_score = score
            best_sheet = sheet_name
    return best_sheet


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit yearlyRecords workbook headers for wording drift.")
    parser.add_argument("folder", nargs="?", default="YearlyRecords", help="Folder containing yearly .xlsx files")
    args = parser.parse_args()

    folder = Path(args.folder)
    files = sorted(folder.glob("*.xlsx"))
    if not files:
        print(f"No .xlsx files found in {folder}")
        return 1

    unknown_total = 0
    for path in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_name = choose_data_sheet(workbook)
        worksheet = workbook[sheet_name]
        headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]

        print(f"FILE: {path.name}  SHEET: {sheet_name}")
        for header in headers:
            if header in (None, ""):
                continue
            if not canonical_header(header):
                unknown_total += 1
                normalized = normalize_header(header)
                suggestions = get_close_matches(normalized, list(HEADER_LOOKUP.keys()), n=3, cutoff=0.72)
                if suggestions:
                    mapped = ", ".join(HEADER_LOOKUP[s] for s in suggestions[:3])
                    print(f"  UNMAPPED: {header!r} -> maybe {mapped}")
                else:
                    print(f"  UNMAPPED: {header!r}")
        print()

    print(f"Unknown headers total: {unknown_total}")
    return 0 if unknown_total == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
